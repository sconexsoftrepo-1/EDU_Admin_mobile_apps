import base64
import io
from decimal import Decimal

from openpyxl import Workbook

from app.core.security import hasher
from app.modules.documents.files.api import pdf


def mobile_login(client, app, kind, person, client_id):
    with app.state.store.transaction() as d:
        d["accounts"][kind] = {
            "code": kind,
            "login_id": kind,
            "password_hash": hasher.hash("Account-Password-123!"),
            "active": True,
            "person_type": kind if kind in {"student", "parent"} else "employee",
            "person_code": person,
            "role_codes": [],
            "clients": [client_id],
        }
    r = client.post(
        "/api/v1/auth/login",
        json={
            "login_id": kind,
            "password": "Account-Password-123!",
            "client": client_id,
        },
    )
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def classroom(app):
    with app.state.store.transaction() as d:
        d["records"].update(
            {
                "parents": {
                    "P1": {"code": "P1", "first_name": "Parent", "last_name": "One"}
                },
                "student_parents": {
                    "SP": {"code": "SP", "student_code": "S1", "parent_code": "P1"}
                },
                "teachers": {"T1": {"code": "T1", "employee_code": "E1"}},
                "subject_offerings": {
                    "OF1": {
                        "code": "OF1",
                        "section_code": "SEC1",
                        "teacher_code": "T1",
                        "subject_code": "SUB",
                    }
                },
                "student_enrollments": {
                    "ENR1": {
                        "code": "ENR1",
                        "student_code": "S1",
                        "section_code": "SEC1",
                        "roll_number": "R1",
                        "status": "active",
                    }
                },
                "certificate_types": {"CT": {"code": "CT", "name": "Bonafide"}},
                "attendance_sessions": {"AS": {"code": "AS", "offering_code": "OF1"}},
                "assignments": {
                    "A1": {
                        "code": "A1",
                        "offering_code": "OF1",
                        "title": "Homework",
                        "instructions": "Read",
                        "due_at": "2099-12-31T00:00:00+00:00",
                        "max_marks": "10",
                    }
                },
                "exams": {},
                "exam_papers": {},
                "leave_types": {
                    "LT": {"code": "LT", "name": "Annual", "annual_days": 12}
                },
            }
        )


def test_every_admin_screen_descriptor(client, headers):
    r = client.get("/api/v1/workspace/screens", headers=headers)
    assert r.status_code == 200, r.text
    assert len(r.json()["screens"]) >= 90
    for s in r.json()["screens"]:
        response = client.get("/api/v1/workspace/screens/" + s["key"], headers=headers)
        assert response.status_code == 200, (s["key"], response.text)


def test_every_mobile_screen_and_ownership(client, seeded):
    classroom(seeded)
    for kind, person, app_id in [
        ("student", "S1", "student-mobile"),
        ("parent", "P1", "parent-mobile"),
        ("teacher", "E1", "teacher-mobile"),
    ]:
        h = mobile_login(client, seeded, kind, person, app_id)
        response = client.get("/api/v1/mobile/workspace/screens", headers=h)
        assert response.status_code == 200
        for s in response.json()["screens"]:
            r = client.get("/api/v1/mobile/workspace/" + s["key"], headers=h)
            assert r.status_code == 200, (app_id, s["key"], r.text)
        assert client.get(
            "/api/v1/mobile/workspace/fees?student_code=S2", headers=h
        ).status_code in {404}


def test_teacher_marks_attendance_and_student_submits(client, seeded):
    classroom(seeded)
    teacher = mobile_login(client, seeded, "teacher", "E1", "teacher-mobile")
    path = "/api/v1/mobile/workspace/attendance/actions/record_attendance"
    assert (
        client.post(
            path,
            headers=teacher,
            json={
                "values": {
                    "session_code": "AS",
                    "student_code": "S1",
                    "status": "present",
                }
            },
        ).status_code
        == 201
    )
    assert (
        client.post(
            path,
            headers=teacher,
            json={
                "values": {
                    "session_code": "AS",
                    "student_code": "S2",
                    "status": "present",
                }
            },
        ).status_code
        == 403
    )
    assert (
        client.post(
            path,
            headers=teacher,
            json={
                "values": {"session_code": "AS", "student_code": "S1", "status": "fake"}
            },
        ).status_code
        == 422
    )
    student = mobile_login(client, seeded, "student", "S1", "student-mobile")
    r = client.post(
        "/api/v1/mobile/workspace/assignments/actions/submit_assignment",
        headers=student,
        json={
            "values": {
                "assignment_code": "A1",
                "document_url": "https://example.org/submission",
            }
        },
    )
    assert r.status_code == 201, r.text
    submission = r.json()["record"]["code"]
    path = "/api/v1/mobile/workspace/submissions/actions/grade_submission"
    assert (
        client.post(
            path,
            headers=teacher,
            json={"values": {"submission_code": submission, "marks": "11"}},
        ).status_code
        == 422
    )
    r = client.post(
        path,
        headers=teacher,
        json={
            "values": {
                "submission_code": submission,
                "marks": "8",
                "feedback": "Good work",
            }
        },
    )
    assert r.status_code == 201, r.text
    assert (
        client.post(
            "/api/v1/mobile/workspace/assignments/actions/submit_assignment",
            headers=student,
            json={
                "values": {
                    "assignment_code": "A1",
                    "document_url": "https://example.org/updated",
                }
            },
        ).status_code
        == 409
    )


def test_payslip_and_form16_download_isolation(client, headers, seeded):
    with seeded.state.store.transaction() as d:
        d["records"]["payroll_runs"] = {"RUN": {"code": "RUN", "status": "published"}}
        d["records"]["payslips"] = {
            "SLIP": {
                "code": "SLIP",
                "employee_code": "E1",
                "payroll_run_code": "RUN",
                "gross_amount": "30000",
                "deduction_amount": "1000",
                "net_amount": "29000",
            }
        }
    teacher = mobile_login(client, seeded, "teacher", "E1", "teacher-mobile")
    r = client.get("/api/v1/documents/payslip/SLIP/download", headers=teacher)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    other = mobile_login(client, seeded, "other", "E2", "teacher-mobile")
    assert (
        client.get("/api/v1/documents/payslip/SLIP/download", headers=other).status_code
        == 404
    )
    upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        json={
            "kind": "form_16",
            "person_code": "E1",
            "filename": "Form16.pdf",
            "financial_year": "2025-2026",
            "content_base64": base64.b64encode(
                pdf("Form 16", ["Issued tax document"])
            ).decode(),
        },
    )
    assert upload.status_code == 201, upload.text
    code = upload.json()["code"]
    assert (
        client.get(
            f"/api/v1/documents/tax-document/{code}/download", headers=teacher
        ).status_code
        == 200
    )
    assert (
        client.get(
            f"/api/v1/documents/tax-document/{code}/download", headers=other
        ).status_code
        == 404
    )


def test_certificate_authorization_signature_and_download_audit(
    client, headers, seeded
):
    classroom(seeded)
    student = mobile_login(client, seeded, "student", "S1", "student-mobile")
    r = client.post(
        "/api/v1/mobile/workspace/certificates/actions/request_certificate",
        headers=student,
        json={"values": {"type_code": "CT", "roll_number": "WRONG", "reason": "Bank"}},
    )
    assert r.status_code == 422
    r = client.post(
        "/api/v1/mobile/workspace/certificates/actions/request_certificate",
        headers=student,
        json={
            "values": {"type_code": "CT", "roll_number": "R1", "reason": "Bank account"}
        },
    )
    assert r.status_code == 201, r.text
    issued = client.post(
        "/api/v1/documents/certificates/issue",
        headers=headers,
        json={"request_code": r.json()["record"]["code"]},
    )
    assert issued.status_code == 201, issued.text
    code = issued.json()["code"]
    state = seeded.state.store.snapshot()
    cert = state["records"]["certificates"][code]
    token = cert["verification_token"]
    assert (
        client.get("/api/v1/documents/certificates/verify/" + token).json()["valid"]
        is True
    )
    assert (
        client.get(
            f"/api/v1/documents/certificate/{code}/download", headers=student
        ).status_code
        == 200
    )
    assert (
        len(seeded.state.store.snapshot()["records"]["certificate_download_events"])
        == 1
    )
    (seeded.state.settings.data_dir / "files" / cert["storage_key"]).write_bytes(
        b"tampered"
    )
    assert (
        client.get("/api/v1/documents/certificates/verify/" + token).json()["valid"]
        is False
    )


def test_atomic_configured_forms(client, headers, app):
    form = client.post(
        "/api/v1/forms",
        headers=headers,
        json={
            "name": "Institution contact",
            "screen": "institutions",
            "published": True,
            "fields": [
                {"key": "contact_person", "label": "Contact person", "required": True}
            ],
        },
    ).json()
    body = {
        "values": {"code": "I1", "name": "College"},
        "form_code": form["code"],
        "form_revision": 1,
    }
    assert (
        client.post(
            "/api/v1/workspace/resources/institutions", headers=headers, json=body
        ).status_code
        == 422
    )
    # The rejected save must not create the institution (SQL runs may hold fixture stub rows).
    assert "I1" not in app.state.store.snapshot()["records"].get("institutions", {})
    body["values"]["contact_person"] = "Example Contact"
    r = client.post(
        "/api/v1/workspace/resources/institutions", headers=headers, json=body
    )
    assert r.status_code == 201, r.text
    assert (
        next(iter(app.state.store.snapshot()["submissions"].values()))["entity_code"]
        == "I1"
    )


def test_stock_purchase_payroll_and_cash_report(client, headers, seeded):
    with seeded.state.store.transaction() as d:
        d["records"].update(
            {
                "inventory_items": {"ITEM": {"code": "ITEM", "quantity": "10"}},
                "payroll_runs": {
                    "RUN": {
                        "code": "RUN",
                        "campus_code": "CAMP",
                        "period_start": "2026-10-01",
                        "period_end": "2026-10-31",
                        "status": "draft",
                    }
                },
                "employee_salary_components": {
                    "SAL": {
                        "code": "SAL",
                        "employee_code": "E1",
                        "component": "Basic",
                        "component_type": "earning",
                        "amount": "30000",
                        "effective_from": "2026-01-01",
                    }
                },
                "expenses": {
                    "EXP": {
                        "code": "EXP",
                        "amount": "100",
                        "payment_status": "unpaid",
                        "expense_date": "2026-10-22",
                    }
                },
            }
        )
        d["records"]["employees"]["E1"].update(campus_code="CAMP", status="active")

    def run(name, v):
        return client.post(
            "/api/v1/workflows/" + name, headers=headers, json={"values": v}
        )

    assert (
        run(
            "stock-move",
            {"item_code": "ITEM", "movement_type": "out", "quantity": "11"},
        ).status_code
        == 409
    )
    assert (
        run(
            "stock-move", {"item_code": "ITEM", "movement_type": "out", "quantity": "2"}
        ).status_code
        == 200
    )
    assert Decimal(
        seeded.state.store.snapshot()["records"]["inventory_items"]["ITEM"]["quantity"]
    ) == Decimal("8")
    assert run("payroll-calculate", {"code": "RUN"}).status_code == 200
    assert run("payroll-calculate", {"code": "RUN"}).status_code == 409
    assert run("payroll-publish", {"code": "RUN"}).status_code == 200
    assert (
        run(
            "expense-pay",
            {
                "code": "EXP",
                "method": "cash",
                "reference": "EXP-REF",
                "paid_on": "2026-10-22",
            },
        ).status_code
        == 200
    )
    assert (
        run(
            "expense-pay",
            {
                "code": "EXP",
                "method": "cash",
                "reference": "EXP-REF",
                "paid_on": "2026-10-22",
            },
        ).status_code
        == 409
    )
    assert Decimal(
        client.get("/api/v1/reports/cashbook.json", headers=headers).json()["items"][
            -1
        ]["balance"]
    ) == Decimal("-100")
    for fmt in ["pdf", "csv", "xlsx"]:
        r = client.get("/api/v1/reports/finance." + fmt, headers=headers)
        assert r.status_code == 200, r.text[:100]


def test_excel_template_preview_and_atomic_commit(client, headers, app):
    from app.modules.imports.excel.api import CONTRACT

    definition = next(
        f for f in CONTRACT["files"] if f["file"] == "12_Leave_Types.xlsx"
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for s in definition["sheets"]:
        ws = workbook.create_sheet(s["name"])
        for _ in range(5):
            ws.append([])
        ws.append([f["label"] + (" *" if f["required"] else "") for f in s["fields"]])
        values = {f["key"]: f["example"] for f in s["fields"]}
        ws.append([values[f["key"]] for f in s["fields"]])
    raw = io.BytesIO()
    workbook.save(raw)
    r = client.post(
        "/api/v1/imports/excel/validate",
        headers=headers,
        json={
            "filename": definition["file"],
            "content_base64": base64.b64encode(raw.getvalue()).decode(),
        },
    )
    assert r.status_code == 201, r.text
    assert r.json()["failed_rows"] == 0, r.text
    code = r.json()["code"]
    r = client.post(f"/api/v1/imports/excel/{code}/commit", headers=headers)
    assert r.status_code == 200, r.text
    assert len(app.state.store.snapshot()["records"]["leave_types"]) == 1
    assert (
        client.post(f"/api/v1/imports/excel/{code}/commit", headers=headers).status_code
        == 409
    )


def test_mobile_custom_form_revision_and_rollback(client, headers, seeded):
    classroom(seeded)
    form = {
        "name": "Employee leave details",
        "screen": "mobile.request_leave",
        "published": True,
        "fields": [
            {"key": "handover_note", "label": "Handover note", "required": True}
        ],
    }
    created = client.post("/api/v1/forms", headers=headers, json=form)
    assert created.status_code == 201, created.text
    teacher = mobile_login(client, seeded, "teacher", "E1", "teacher-mobile")
    path = "/api/v1/mobile/workspace/leave/actions/request_leave"
    values = {
        "leave_type_code": "LT",
        "start_date": "2026-11-01",
        "end_date": "2026-11-02",
        "reason": "Family event",
    }
    assert (
        client.post(path, headers=teacher, json={"values": values}).status_code == 409
    )
    assert (
        client.post(
            path, headers=teacher, json={"values": values, "form_revision": 1}
        ).status_code
        == 422
    )
    assert not seeded.state.store.snapshot()["records"].get("leave_requests")
    response = client.post(
        path,
        headers=teacher,
        json={
            "values": {**values, "handover_note": "Covered by colleague"},
            "form_revision": 1,
        },
    )
    assert response.status_code == 201, response.text
    assert (
        next(iter(seeded.state.store.snapshot()["submissions"].values()))["entity_code"]
        == response.json()["record"]["code"]
    )
    code = created.json()["code"]
    assert (
        client.put(
            "/api/v1/forms/" + code, headers=headers, json={**form, "published": False}
        ).status_code
        == 200
    )
    descriptor = client.get("/api/v1/mobile/workspace/leave", headers=teacher).json()
    assert "handover_note" not in {f["key"] for f in descriptor["actions"][0]["fields"]}


def test_teaching_scope_web_cannot_use_global_student_endpoint(client, seeded):
    classroom(seeded)
    h = mobile_login(client, seeded, "teacher", "E1", "teacher-mobile")
    with seeded.state.store.transaction() as data:
        data["roles"]["FAC"] = {
            "code": "FAC",
            "name": "Faculty",
            "scope": "teaching",
            "permissions": {"students": ["view", "edit"]},
        }
        data["accounts"]["teacher"]["role_codes"] = ["FAC"]
        data["accounts"]["teacher"]["clients"].append("admin-web")
    r = client.post(
        "/api/v1/auth/login",
        json={
            "login_id": "teacher",
            "password": "Account-Password-123!",
            "client": "admin-web",
        },
    )
    h = {"Authorization": "Bearer " + r.json()["access_token"]}
    assert (
        client.get("/api/v1/auth/me", headers=h).json()["account"]["workspace_scope"]
        == "teaching"
    )
    assert client.get("/api/v1/mobile/workspace/classes", headers=h).status_code == 200
    assert (
        client.get("/api/v1/student-management/students", headers=h).status_code == 403
    )


def test_scoped_message_thread_and_parent_performance(client, seeded):
    classroom(seeded)
    parent = mobile_login(client, seeded, "parent", "P1", "parent-mobile")
    teacher = mobile_login(client, seeded, "teacher", "E1", "teacher-mobile")
    endpoint = "/api/v1/mobile/workspace/messages/actions/start_thread"
    values = {
        "recipient_user_code": "teacher",
        "subject": "Homework",
        "body": "Please clarify the assignment",
    }
    assert (
        client.post(
            endpoint,
            headers=parent,
            json={"values": {**values, "recipient_user_code": "ADMIN"}},
        ).status_code
        == 403
    )
    r = client.post(endpoint, headers=parent, json={"values": values})
    assert r.status_code == 201, r.text
    thread = r.json()["record"]["thread_code"]
    r = client.post(
        "/api/v1/mobile/workspace/messages/actions/send_message",
        headers=teacher,
        json={"values": {"thread_code": thread, "body": "We will discuss it tomorrow"}},
    )
    assert r.status_code == 201, r.text
    assert (
        client.get("/api/v1/mobile/workspace/messages", headers=parent).json()["total"]
        == 2
    )


def test_excel_recognizes_renamed_templates_without_bypassing_role_limits(client, headers, app):
    from app.modules.imports.excel.api import CONTRACT
    definition = next(f for f in CONTRACT['files'] if f['file'] == '12_Leave_Types.xlsx')
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in definition['sheets']:
        ws = workbook.create_sheet(sheet['name'])
        for _ in range(5):
            ws.append([])
        ws.append([field['label'] + (' *' if field['required'] else '') for field in sheet['fields']])
        ws.append([field['example'] for field in sheet['fields']])
    raw = io.BytesIO()
    workbook.save(raw)
    payload = {'filename':'Leave Types - completed (1).xlsx','content_base64':base64.b64encode(raw.getvalue()).decode()}
    response = client.post('/api/v1/imports/excel/validate', headers=headers, json=payload)
    assert response.status_code == 201, response.text
    assert response.json()['filename'] == definition['file']
    assert response.json()['failed_rows'] == 0
    committed = client.post('/api/v1/imports/excel/' + response.json()['code'] + '/commit', headers=headers)
    assert committed.status_code == 200, committed.text
    # Detection must use the actual workbook contract for the role allowlist.
    with app.state.store.transaction() as data:
        data['roles']['ROOT']['system_admin'] = False
        data['roles']['ROOT']['scope'] = 'college'
        data['roles']['ROOT']['permissions'] = {'import_jobs':['view','create','approve']}
        data['roles']['ROOT']['import_files'] = ['01_Students_and_Parents.xlsx']
    denied = client.post('/api/v1/imports/excel/validate', headers=headers, json=payload)
    assert denied.status_code == 403


def test_excel_rejects_unrecognized_workbooks_with_an_actionable_message(client, headers):
    workbook = Workbook()
    workbook.active.append(['Unrelated spreadsheet'])
    raw = io.BytesIO()
    workbook.save(raw)
    response = client.post('/api/v1/imports/excel/validate', headers=headers,
        json={'filename':'notes.xlsx','content_base64':base64.b64encode(raw.getvalue()).decode()})
    assert response.status_code == 422
    assert 'sheet names and headers in row 6' in response.json()['detail']
